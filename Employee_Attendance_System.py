import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os
from datetime import datetime


# =========================================================
# PROFESSIONAL COLOR THEME
# =========================================================

# Navy + blue is used for a clean, professional business application.
BG_COLOR = "#F4F7FB"
CARD_COLOR = "#FFFFFF"
PRIMARY_COLOR = "#1F4E78"
PRIMARY_HOVER = "#173A5B"
ACCENT_COLOR = "#2F75B5"
TEXT_COLOR = "#1F2937"
MUTED_COLOR = "#6B7280"
SUCCESS_COLOR = "#2E7D32"
DANGER_COLOR = "#C62828"
BORDER_COLOR = "#D9E2EC"


# =========================================================
# EMPLOYEE ATTENDANCE MANAGEMENT SYSTEM
# Python + Tkinter + Excel
# =========================================================

# Excel file will be created in the same folder as this Python file
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FILE_NAME = os.path.join(BASE_DIR, "employee_attendance.xlsx")

SHEET_NAME = "Attendance"

COLUMNS = [
    "Record ID",
    "Employee ID",
    "Employee Name",
    "Department",
    "Date",
    "Status",
    "Check In",
    "Check Out"
]


# =========================================================
# EXCEL FUNCTIONS
# =========================================================

def create_excel_file():
    """Create the Excel file and header row if it does not exist."""
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(COLUMNS)
        workbook.save(FILE_NAME)


def get_workbook():
    """Open and return the attendance workbook."""
    return load_workbook(FILE_NAME)


def get_next_record_id(sheet):
    """Generate the next Record ID."""
    record_ids = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            try:
                record_ids.append(int(row[0]))
            except (ValueError, TypeError):
                pass

    return max(record_ids, default=0) + 1


def get_all_records():
    """Read all attendance records from Excel."""
    workbook = get_workbook()
    sheet = workbook[SHEET_NAME]

    records = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(value is not None for value in row):
            records.append(list(row))

    workbook.close()
    return records


def find_row_by_record_id(record_id):
    """Find the Excel row number for a Record ID."""
    workbook = get_workbook()
    sheet = workbook[SHEET_NAME]

    for row_number in range(2, sheet.max_row + 1):
        value = sheet.cell(row=row_number, column=1).value

        if str(value) == str(record_id):
            workbook.close()
            return row_number

    workbook.close()
    return None


def attendance_exists(employee_id, date, ignore_record_id=None):
    """Prevent duplicate attendance for the same employee on the same date."""
    records = get_all_records()

    for record in records:
        record_id = str(record[0])
        existing_employee_id = str(record[1]).strip().lower()
        existing_date = str(record[4]).strip()

        if ignore_record_id is not None and record_id == str(ignore_record_id):
            continue

        if (
            existing_employee_id == employee_id.strip().lower()
            and existing_date == date.strip()
        ):
            return True

    return False


# =========================================================
# VALIDATION
# =========================================================

def validate_attendance(employee_id, employee_name, department,
                        date, status, check_in, check_out,
                        ignore_record_id=None):

    if not employee_id.strip():
        messagebox.showwarning("Validation", "Please enter Employee ID.")
        return False

    if not employee_name.strip():
        messagebox.showwarning("Validation", "Please enter Employee Name.")
        return False

    if not department.strip():
        messagebox.showwarning("Validation", "Please enter Department.")
        return False

    if not date.strip():
        messagebox.showwarning("Validation", "Please enter Date.")
        return False

    # Check date format
    try:
        datetime.strptime(date.strip(), "%d-%m-%Y")
    except ValueError:
        messagebox.showwarning(
            "Validation",
            "Date must be in DD-MM-YYYY format."
        )
        return False

    if not status:
        messagebox.showwarning("Validation", "Please select Attendance Status.")
        return False

    # Check time format only when a time is entered
    for time_value, label in [
        (check_in, "Check In"),
        (check_out, "Check Out")
    ]:
        if time_value.strip():
            try:
                datetime.strptime(time_value.strip(), "%I:%M %p")
            except ValueError:
                messagebox.showwarning(
                    "Validation",
                    f"{label} must be in format HH:MM AM/PM.\n"
                    f"Example: 09:30 AM"
                )
                return False

    if attendance_exists(
        employee_id,
        date,
        ignore_record_id
    ):
        messagebox.showwarning(
            "Duplicate Record",
            "Attendance for this Employee ID on this date already exists."
        )
        return False

    return True


# =========================================================
# MAIN APPLICATION
# =========================================================

class AttendanceApp:

    def __init__(self, root):
        self.root = root
        self.root.title("Employee Attendance Management System")
        self.root.geometry("950x650")
        self.root.resizable(False, False)
        self.root.configure(bg=BG_COLOR)

        self.selected_record_id = None

        self.setup_style()
        self.create_login_page()

    # =====================================================
    # STYLE
    # =====================================================

    def setup_style(self):
        style = ttk.Style()

        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        # Main application background
        self.root.configure(bg=BG_COLOR)

        style.configure(
            "TFrame",
            background=BG_COLOR
        )

        style.configure(
            "Card.TFrame",
            background=CARD_COLOR
        )

        style.configure(
            "Title.TLabel",
            background=BG_COLOR,
            foreground=PRIMARY_COLOR,
            font=("Segoe UI", 24, "bold")
        )

        style.configure(
            "Subtitle.TLabel",
            background=BG_COLOR,
            foreground=MUTED_COLOR,
            font=("Segoe UI", 13)
        )

        style.configure(
            "Form.TLabel",
            background=CARD_COLOR,
            foreground=TEXT_COLOR,
            font=("Segoe UI", 11)
        )

        style.configure(
            "Hint.TLabel",
            background=BG_COLOR,
            foreground=MUTED_COLOR,
            font=("Segoe UI", 10)
        )

        style.configure(
            "TButton",
            font=("Segoe UI", 10, "bold"),
            padding=(16, 9),
            foreground=TEXT_COLOR
        )

        style.configure(
            "Primary.TButton",
            font=("Segoe UI", 10, "bold"),
            padding=(18, 10),
            foreground="white",
            background=PRIMARY_COLOR
        )

        style.map(
            "Primary.TButton",
            background=[
                ("active", PRIMARY_HOVER),
                ("pressed", PRIMARY_HOVER)
            ],
            foreground=[
                ("active", "white"),
                ("pressed", "white")
            ]
        )

        style.configure(
            "Danger.TButton",
            font=("Segoe UI", 10, "bold"),
            padding=(18, 10),
            foreground="white",
            background=DANGER_COLOR
        )

        style.map(
            "Danger.TButton",
            background=[
                ("active", "#9B1C1C"),
                ("pressed", "#9B1C1C")
            ],
            foreground=[
                ("active", "white"),
                ("pressed", "white")
            ]
        )

        style.configure(
            "Treeview.Heading",
            background=PRIMARY_COLOR,
            foreground="white",
            font=("Segoe UI", 10, "bold"),
            padding=8
        )

        style.map(
            "Treeview.Heading",
            background=[("active", PRIMARY_HOVER)]
        )

        style.configure(
            "Treeview",
            background=CARD_COLOR,
            fieldbackground=CARD_COLOR,
            foreground=TEXT_COLOR,
            font=("Segoe UI", 10),
            rowheight=30
        )

        style.map(
            "Treeview",
            background=[("selected", ACCENT_COLOR)],
            foreground=[("selected", "white")]
        )

        style.configure(
            "TEntry",
            padding=7,
            font=("Segoe UI", 10)
        )

        style.configure(
            "TCombobox",
            padding=6,
            font=("Segoe UI", 10)
        )

    # =====================================================
    # LOGIN PAGE
    # =====================================================

    def create_login_page(self):
        self.clear_root()

        self.login_frame = ttk.Frame(
            self.root,
            style="Card.TFrame",
            padding=(45, 35)
        )
        self.login_frame.pack(expand=True)

        ttk.Label(
            self.login_frame,
            text="EMPLOYEE ATTENDANCE",
            font=("Segoe UI", 23, "bold"),
            foreground=PRIMARY_COLOR,
            background=CARD_COLOR
        ).pack(pady=(10, 3))

        ttk.Label(
            self.login_frame,
            text="Management System",
            font=("Segoe UI", 14),
            foreground=MUTED_COLOR,
            background=CARD_COLOR
        ).pack(pady=(0, 28))

        form = ttk.Frame(
            self.login_frame,
            style="Card.TFrame"
        )
        form.pack()

        ttk.Label(
            form,
            text="Username:",
            font=("Segoe UI", 11)
        ).grid(row=0, column=0, padx=10, pady=10, sticky="w")

        self.username_entry = ttk.Entry(form, width=30)
        self.username_entry.grid(row=0, column=1, padx=10, pady=10)

        ttk.Label(
            form,
            text="Password:",
            font=("Segoe UI", 11)
        ).grid(row=1, column=0, padx=10, pady=10, sticky="w")

        self.password_entry = ttk.Entry(
            form,
            width=30,
            show="*"
        )
        self.password_entry.grid(row=1, column=1, padx=10, pady=10)

        button_frame = ttk.Frame(
            self.login_frame,
            style="Card.TFrame"
        )
        button_frame.pack(pady=25)

        ttk.Button(
            button_frame,
            text="LOGIN",
            width=15,
            style="Primary.TButton",
            command=self.login
        ).grid(row=0, column=0, padx=8)

        ttk.Button(
            button_frame,
            text="CLEAR",
            width=15,
            command=self.clear_login
        ).grid(row=0, column=1, padx=8)

        ttk.Label(
            self.login_frame,
            text="Demo Login  •  Username: admin  |  Password: admin123",
            style="Hint.TLabel"
        ).pack(pady=15)

        self.username_entry.focus()

        self.root.bind("<Return>", lambda event: self.login())

    def login(self):
        username = self.username_entry.get().strip()
        password = self.password_entry.get()

        if username == "admin" and password == "admin123":
            self.root.unbind("<Return>")
            self.show_dashboard()
        else:
            messagebox.showerror(
                "Login Error",
                "Invalid username or password."
            )

    def clear_login(self):
        self.username_entry.delete(0, tk.END)
        self.password_entry.delete(0, tk.END)
        self.username_entry.focus()

    # =====================================================
    # DASHBOARD
    # =====================================================

    def show_dashboard(self):
        self.clear_root()

        self.dashboard_frame = ttk.Frame(
            self.root,
            style="TFrame"
        )
        self.dashboard_frame.pack(fill="both", expand=True)

        ttk.Label(
            self.dashboard_frame,
            text="EMPLOYEE ATTENDANCE MANAGEMENT SYSTEM",
            font=("Segoe UI", 23, "bold"),
            foreground=PRIMARY_COLOR,
            background=BG_COLOR
        ).pack(pady=(30, 5))

        ttk.Label(
            self.dashboard_frame,
            text="Attendance Management Dashboard",
            style="Subtitle.TLabel"
        ).pack(pady=(0, 30))

        button_frame = ttk.Frame(self.dashboard_frame)
        button_frame.pack()

        buttons = [
            ("➕  Add Attendance", self.open_add_window),
            ("📋  View Records", self.open_view_window),
            ("🔎  Search Records", self.open_search_window),
            ("✏  Update Record", self.open_update_window),
            ("🗑  Delete Record", self.open_delete_window),
            ("↻  Clear / Reset", self.open_add_window)
        ]

        for index, (text, command) in enumerate(buttons):
            row = index // 2
            column = index % 2

            ttk.Button(
                button_frame,
                text=text,
                width=25,
                style="Primary.TButton",
                command=command
            ).grid(
                row=row,
                column=column,
                padx=20,
                pady=15
            )

        ttk.Button(
            self.dashboard_frame,
            text="LOGOUT",
            width=20,
            style="Danger.TButton",
            command=self.logout
        ).pack(pady=25)

        ttk.Label(
            self.dashboard_frame,
            text="Python  •  Tkinter  •  Excel",
            style="Hint.TLabel"
        ).pack(pady=(0, 10))

    # =====================================================
    # ADD / UPDATE FORM
    # =====================================================

    def open_add_window(self):
        self.open_attendance_form(mode="add")

    def open_attendance_form(self, mode="add", record=None):
        form_window = tk.Toplevel(self.root)

        if mode == "add":
            form_window.title("Add Attendance")
        else:
            form_window.title("Update Attendance")

        form_window.geometry("560x600")
        form_window.resizable(False, False)
        form_window.configure(bg=BG_COLOR)
        form_window.grab_set()

        title_text = (
            "Add Employee Attendance"
            if mode == "add"
            else "Update Employee Attendance"
        )

        ttk.Label(
            form_window,
            text=title_text,
            font=("Segoe UI", 20, "bold")
        ).pack(pady=20)

        form = ttk.Frame(form_window)
        form.pack(pady=5)

        ttk.Label(
            form,
            text="Employee ID:",
            style="Form.TLabel"
        ).grid(row=0, column=0, padx=10, pady=10, sticky="w")

        employee_id_entry = ttk.Entry(form, width=32)
        employee_id_entry.grid(row=0, column=1, padx=10, pady=10)

        ttk.Label(
            form,
            text="Employee Name:",
            style="Form.TLabel"
        ).grid(row=1, column=0, padx=10, pady=10, sticky="w")

        employee_name_entry = ttk.Entry(form, width=32)
        employee_name_entry.grid(row=1, column=1, padx=10, pady=10)

        ttk.Label(
            form,
            text="Department:",
            style="Form.TLabel"
        ).grid(row=2, column=0, padx=10, pady=10, sticky="w")

        department_entry = ttk.Entry(form, width=32)
        department_entry.grid(row=2, column=1, padx=10, pady=10)

        ttk.Label(
            form,
            text="Date (DD-MM-YYYY):",
            style="Form.TLabel"
        ).grid(row=3, column=0, padx=10, pady=10, sticky="w")

        date_entry = ttk.Entry(form, width=32)
        date_entry.grid(row=3, column=1, padx=10, pady=10)

        ttk.Label(
            form,
            text="Status:",
            style="Form.TLabel"
        ).grid(row=4, column=0, padx=10, pady=10, sticky="w")

        status_combo = ttk.Combobox(
            form,
            values=["Present", "Absent", "Leave"],
            width=29,
            state="readonly"
        )
        status_combo.grid(row=4, column=1, padx=10, pady=10)

        ttk.Label(
            form,
            text="Check In:",
            style="Form.TLabel"
        ).grid(row=5, column=0, padx=10, pady=10, sticky="w")

        checkin_entry = ttk.Entry(form, width=32)
        checkin_entry.grid(row=5, column=1, padx=10, pady=10)

        ttk.Label(
            form,
            text="Check Out:",
            style="Form.TLabel"
        ).grid(row=6, column=0, padx=10, pady=10, sticky="w")

        checkout_entry = ttk.Entry(form, width=32)
        checkout_entry.grid(row=6, column=1, padx=10, pady=10)

        # Fill form when updating
        if mode == "update" and record:
            employee_id_entry.insert(0, record[1] or "")
            employee_name_entry.insert(0, record[2] or "")
            department_entry.insert(0, record[3] or "")
            date_entry.insert(0, record[4] or "")
            status_combo.set(record[5] or "Present")
            checkin_entry.insert(0, record[6] or "")
            checkout_entry.insert(0, record[7] or "")
        else:
            date_entry.insert(
                0,
                datetime.now().strftime("%d-%m-%Y")
            )
            status_combo.set("Present")

        button_frame = ttk.Frame(form_window)
        button_frame.pack(pady=30)

        def clear_fields():
            employee_id_entry.delete(0, tk.END)
            employee_name_entry.delete(0, tk.END)
            department_entry.delete(0, tk.END)
            date_entry.delete(0, tk.END)
            checkin_entry.delete(0, tk.END)
            checkout_entry.delete(0, tk.END)

            date_entry.insert(
                0,
                datetime.now().strftime("%d-%m-%Y")
            )
            status_combo.set("Present")
            employee_id_entry.focus()

        def save():
            employee_id = employee_id_entry.get()
            employee_name = employee_name_entry.get()
            department = department_entry.get()
            date = date_entry.get()
            status = status_combo.get()
            check_in = checkin_entry.get()
            check_out = checkout_entry.get()

            ignore_id = record[0] if mode == "update" and record else None

            if not validate_attendance(
                employee_id,
                employee_name,
                department,
                date,
                status,
                check_in,
                check_out,
                ignore_id
            ):
                return

            workbook = get_workbook()
            sheet = workbook[SHEET_NAME]

            if mode == "add":
                record_id = get_next_record_id(sheet)

                sheet.append([
                    record_id,
                    employee_id.strip(),
                    employee_name.strip(),
                    department.strip(),
                    date.strip(),
                    status,
                    check_in.strip(),
                    check_out.strip()
                ])

                message = "Attendance record added successfully!"

            else:
                row_number = find_row_by_record_id(record[0])

                if row_number is None:
                    workbook.close()
                    messagebox.showerror(
                        "Update Error",
                        "Record could not be found."
                    )
                    return

                values = [
                    employee_id.strip(),
                    employee_name.strip(),
                    department.strip(),
                    date.strip(),
                    status,
                    check_in.strip(),
                    check_out.strip()
                ]

                for column_number, value in enumerate(values, start=2):
                    sheet.cell(
                        row=row_number,
                        column=column_number
                    ).value = value

                message = "Attendance record updated successfully!"

            workbook.save(FILE_NAME)
            workbook.close()

            messagebox.showinfo("Success", message)
            form_window.destroy()

        ttk.Button(
            button_frame,
            text="SAVE ATTENDANCE" if mode == "add" else "UPDATE RECORD",
            width=18,
            style="Primary.TButton",
            command=save
        ).grid(row=0, column=0, padx=10)

        ttk.Button(
            button_frame,
            text="Clear / Reset",
            width=18,
            command=clear_fields
        ).grid(row=0, column=1, padx=10)

        ttk.Button(
            button_frame,
            text="Close",
            width=18,
            command=form_window.destroy
        ).grid(row=1, column=0, columnspan=2, pady=15)

    # =====================================================
    # TREEVIEW HELPER
    # =====================================================

    def create_treeview(self, parent):
        frame = ttk.Frame(parent)
        frame.pack(fill="both", expand=True, padx=10, pady=10)

        tree = ttk.Treeview(
            frame,
            columns=COLUMNS,
            show="headings",
            selectmode="browse"
        )

        widths = {
            "Record ID": 70,
            "Employee ID": 90,
            "Employee Name": 130,
            "Department": 100,
            "Date": 100,
            "Status": 90,
            "Check In": 90,
            "Check Out": 90
        }

        for column in COLUMNS:
            tree.heading(column, text=column)
            tree.column(
                column,
                width=widths.get(column, 100),
                anchor="center"
            )

        scrollbar_y = ttk.Scrollbar(
            frame,
            orient="vertical",
            command=tree.yview
        )

        scrollbar_x = ttk.Scrollbar(
            frame,
            orient="horizontal",
            command=tree.xview
        )

        tree.configure(
            yscrollcommand=scrollbar_y.set,
            xscrollcommand=scrollbar_x.set
        )

        tree.grid(row=0, column=0, sticky="nsew")
        scrollbar_y.grid(row=0, column=1, sticky="ns")
        scrollbar_x.grid(row=1, column=0, sticky="ew")

        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)

        return tree

    def fill_treeview(self, tree, records=None):
        for item in tree.get_children():
            tree.delete(item)

        if records is None:
            records = get_all_records()

        for record in records:
            tree.insert(
                "",
                tk.END,
                values=record
            )

    def get_selected_record(self, tree):
        selection = tree.selection()

        if not selection:
            messagebox.showwarning(
                "Selection Required",
                "Please select a record first."
            )
            return None

        values = tree.item(selection[0], "values")

        return list(values)

    # =====================================================
    # VIEW RECORDS
    # =====================================================

    def open_view_window(self):
        view_window = tk.Toplevel(self.root)
        view_window.title("View Attendance Records")
        view_window.geometry("1050x550")
        view_window.resizable(False, False)
        view_window.configure(bg=BG_COLOR)

        ttk.Label(
            view_window,
            text="Attendance Records",
            font=("Segoe UI", 20, "bold")
        ).pack(pady=15)

        tree = self.create_treeview(view_window)
        self.fill_treeview(tree)

        button_frame = ttk.Frame(view_window)
        button_frame.pack(pady=10)

        def update_selected():
            record = self.get_selected_record(tree)

            if record:
                self.open_attendance_form(
                    mode="update",
                    record=record
                )

        def delete_selected():
            record = self.get_selected_record(tree)

            if record:
                self.delete_record_by_id(
                    record[0],
                    refresh_callback=lambda: self.fill_treeview(tree)
                )

        ttk.Button(
            button_frame,
            text="Update Selected",
            width=18,
            command=update_selected
        ).grid(row=0, column=0, padx=8)

        ttk.Button(
            button_frame,
            text="Delete Selected",
            width=18,
            style="Danger.TButton",
            command=delete_selected
        ).grid(row=0, column=1, padx=8)

        ttk.Button(
            button_frame,
            text="Refresh",
            width=15,
            command=lambda: self.fill_treeview(tree)
        ).grid(row=0, column=2, padx=8)

        ttk.Button(
            button_frame,
            text="Close",
            width=15,
            command=view_window.destroy
        ).grid(row=0, column=3, padx=8)

    # =====================================================
    # SEARCH RECORDS
    # =====================================================

    def open_search_window(self):
        search_window = tk.Toplevel(self.root)
        search_window.title("Search Attendance Records")
        search_window.geometry("1050x600")
        search_window.resizable(False, False)
        search_window.configure(bg=BG_COLOR)

        ttk.Label(
            search_window,
            text="Search Attendance Records",
            font=("Segoe UI", 20, "bold")
        ).pack(pady=15)

        search_frame = ttk.Frame(search_window)
        search_frame.pack(pady=5)

        ttk.Label(
            search_frame,
            text="Employee ID / Name:"
        ).grid(row=0, column=0, padx=10)

        search_entry = ttk.Entry(
            search_frame,
            width=35
        )
        search_entry.grid(row=0, column=1, padx=10)

        tree = self.create_treeview(search_window)

        def search():
            keyword = search_entry.get().strip().lower()

            if not keyword:
                self.fill_treeview(tree)
                return

            records = []

            for record in get_all_records():
                employee_id = str(record[1] or "").lower()
                employee_name = str(record[2] or "").lower()

                if (
                    keyword in employee_id
                    or keyword in employee_name
                ):
                    records.append(record)

            self.fill_treeview(tree, records)

            if not records:
                messagebox.showinfo(
                    "Search Result",
                    "No matching records found."
                )

        def clear_search():
            search_entry.delete(0, tk.END)
            self.fill_treeview(tree)
            search_entry.focus()

        def update_selected():
            record = self.get_selected_record(tree)

            if record:
                self.open_attendance_form(
                    mode="update",
                    record=record
                )

        def delete_selected():
            record = self.get_selected_record(tree)

            if record:
                self.delete_record_by_id(
                    record[0],
                    refresh_callback=lambda: search()
                )

        ttk.Button(
            search_frame,
            text="Search",
            width=15,
            command=search
        ).grid(row=0, column=2, padx=5)

        ttk.Button(
            search_frame,
            text="Clear",
            width=15,
            command=clear_search
        ).grid(row=0, column=3, padx=5)

        button_frame = ttk.Frame(search_window)
        button_frame.pack(pady=10)

        ttk.Button(
            button_frame,
            text="Update Selected",
            width=18,
            command=update_selected
        ).grid(row=0, column=0, padx=8)

        ttk.Button(
            button_frame,
            text="Delete Selected",
            width=18,
            style="Danger.TButton",
            command=delete_selected
        ).grid(row=0, column=1, padx=8)

        ttk.Button(
            button_frame,
            text="Close",
            width=15,
            command=search_window.destroy
        ).grid(row=0, column=2, padx=8)

        search_entry.focus()

    # =====================================================
    # UPDATE RECORD
    # =====================================================

    def open_update_window(self):
        update_window = tk.Toplevel(self.root)
        update_window.title("Select Record to Update")
        update_window.geometry("1050x550")
        update_window.resizable(False, False)
        update_window.configure(bg=BG_COLOR)

        ttk.Label(
            update_window,
            text="Select Attendance Record to Update",
            font=("Segoe UI", 20, "bold")
        ).pack(pady=15)

        tree = self.create_treeview(update_window)
        self.fill_treeview(tree)

        def update_selected():
            record = self.get_selected_record(tree)

            if record:
                update_window.destroy()

                self.open_attendance_form(
                    mode="update",
                    record=record
                )

        ttk.Button(
            update_window,
            text="Update Selected Record",
            width=25,
            command=update_selected
        ).pack(pady=10)

        ttk.Button(
            update_window,
            text="Close",
            width=15,
            command=update_window.destroy
        ).pack(pady=5)

    # =====================================================
    # DELETE RECORD
    # =====================================================

    def open_delete_window(self):
        delete_window = tk.Toplevel(self.root)
        delete_window.title("Select Record to Delete")
        delete_window.geometry("1050x550")
        delete_window.resizable(False, False)
        delete_window.configure(bg=BG_COLOR)

        ttk.Label(
            delete_window,
            text="Select Attendance Record to Delete",
            font=("Segoe UI", 20, "bold")
        ).pack(pady=15)

        tree = self.create_treeview(delete_window)
        self.fill_treeview(tree)

        def delete_selected():
            record = self.get_selected_record(tree)

            if record:
                self.delete_record_by_id(
                    record[0],
                    refresh_callback=lambda: self.fill_treeview(tree)
                )

        ttk.Button(
            delete_window,
            text="Delete Selected Record",
            width=25,
            style="Danger.TButton",
            command=delete_selected
        ).pack(pady=10)

        ttk.Button(
            delete_window,
            text="Refresh",
            width=15,
            command=lambda: self.fill_treeview(tree)
        ).pack(pady=5)

        ttk.Button(
            delete_window,
            text="Close",
            width=15,
            command=delete_window.destroy
        ).pack(pady=5)

    def delete_record_by_id(self, record_id, refresh_callback=None):
        confirm = messagebox.askyesno(
            "Confirm Deletion",
            f"Are you sure you want to delete Record ID {record_id}?"
        )

        if not confirm:
            return

        workbook = get_workbook()
        sheet = workbook[SHEET_NAME]

        row_number = None

        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row=row, column=1).value) == str(record_id):
                row_number = row
                break

        if row_number is None:
            workbook.close()

            messagebox.showerror(
                "Delete Error",
                "Record could not be found."
            )
            return

        sheet.delete_rows(row_number, 1)

        workbook.save(FILE_NAME)
        workbook.close()

        messagebox.showinfo(
            "Success",
            "Attendance record deleted successfully!"
        )

        if refresh_callback:
            refresh_callback()

    # =====================================================
    # LOGOUT
    # =====================================================

    def logout(self):
        confirm = messagebox.askyesno(
            "Logout",
            "Are you sure you want to logout?"
        )

        if confirm:
            self.show_login()

    def show_login(self):
        self.clear_root()
        self.create_login_page()

    # =====================================================
    # CLEAR ROOT WINDOW
    # =====================================================

    def clear_root(self):
        for widget in self.root.winfo_children():
            widget.destroy()


# =========================================================
# PROGRAM START
# =========================================================

if __name__ == "__main__":

    # Create Excel storage file
    create_excel_file()

    # Create main Tkinter window
    root = tk.Tk()

    # Start application
    app = AttendanceApp(root)

    # Run Tkinter application
    root.mainloop()
